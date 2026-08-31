from pathlib import Path
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest
from sqlalchemy import event, func, select
from sqlalchemy.orm import Session

from app.models.budget_entry import BudgetEntry
from app.repositories.budget_repository import BudgetRepository
from finance_server.app_factory import create_app
from finance_server.config import ServerSettings
from finance_server.models import AuditLog, User, UserRole
from finance_server.security import hash_password


PASSWORD = "Strong!Pass123"


def _budget_workbook(path: Path, *, duplicate: bool = False) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planej. orçamentário"
    for _ in range(8):
        sheet.append([])
    sheet.append([None, "Repasse CDLs Estado Goiás", 200, None, 250])
    sheet.append([None, "Despesas com Pessoal", 100, None, 120])
    if duplicate:
        sheet.append([None, "Despesas com Pessoal", 100])
    workbook.save(path)
    return path


def _target_workbook(path: Path) -> Path:
    workbook = Workbook()
    target = workbook.active
    target.title = "Meta"
    target.append(["COD.", "Entidade", "JAN"])
    target.append([7001, "Entidade", 10])
    actual = workbook.create_sheet("Faturamento")
    actual.append(["COD.", "Entidade", "JAN"])
    actual.append([7001, "Entidade", 8])
    workbook.save(path)
    return path


@pytest.fixture
def budget_import_context(tmp_path):
    settings = ServerSettings(
        f"sqlite:///{(tmp_path / 'server.db').as_posix()}", "s" * 64
    )
    app = create_app(settings, create_schema=True)
    with app.state.session_factory() as db:
        db.add_all([
            User(
                nome="Admin", email="admin@test.local", username="admin",
                password_hash=hash_password(PASSWORD),
                perfil=UserRole.ADMINISTRATOR.value, ativo=True,
            ),
            User(
                nome="Gestor", email="gestor@test.local", username="gestor",
                password_hash=hash_password(PASSWORD),
                perfil=UserRole.MANAGER.value, ativo=True,
            ),
            User(
                nome="Consulta", email="consulta@test.local", username="consulta",
                password_hash=hash_password(PASSWORD),
                perfil=UserRole.READ_ONLY.value, ativo=True,
            ),
        ])
        db.commit()
    with TestClient(app) as client:
        yield client, app, tmp_path
    app.state.engine.dispose()


def _headers(client, username="admin"):
    pair = client.post("/api/v1/auth/login", json={
        "identifier": username, "password": PASSWORD,
    }).json()
    return {"Authorization": f"Bearer {pair['access_token']}"}


def _upload(client, endpoint, path, headers):
    with path.open("rb") as handle:
        return client.post(
            endpoint, headers=headers,
            files={"file": (path.name, handle)},
        )


def _seed_budgets(app):
    with app.state.session_factory() as db:
        first = BudgetEntry(
            periodo_ano=2026,
            periodo_mes=8,
            tipo="DESPESA",
            categoria="DIRETORIA",
            descricao="Alimentação/Refeição",
            valor_orcado=2,
            observacao="TESTE BUG-ORC-002",
        )
        second = BudgetEntry(
            periodo_ano=2026,
            periodo_mes=9,
            tipo="DESPESA",
            categoria="DIRETORIA",
            descricao="Alimentação/Refeição",
            valor_orcado=3,
            observacao="TESTE PRESERVADO",
        )
        db.add_all([first, second])
        db.commit()
        return first.id, second.id


def test_administrator_deletes_only_requested_budget_and_audits(
    budget_import_context,
):
    client, app, _ = budget_import_context
    first_id, second_id = _seed_budgets(app)

    response = client.delete(
        f"/api/v1/budgets/{first_id}", headers=_headers(client, "admin")
    )

    assert response.status_code == 204
    assert response.content == b""
    with app.state.session_factory() as db:
        assert db.get(BudgetEntry, first_id) is None
        assert db.get(BudgetEntry, second_id) is not None
        deletion = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "BUDGET_DELETED",
                AuditLog.entity_id == str(first_id),
            )
        )
        assert deletion is not None
        assert deletion.entity_type == "BudgetEntry"
        assert deletion.details == {
            "year": 2026,
            "month": 8,
            "type": "DESPESA",
            "category": "DIRETORIA",
            "description": "Alimentação/Refeição",
            "budgeted_value": "2.0000",
            "notes": "TESTE BUG-ORC-002",
        }
        assert json.loads(json.dumps(deletion.details)) == deletion.details


def test_manager_can_delete_budget(budget_import_context):
    client, app, _ = budget_import_context
    first_id, _ = _seed_budgets(app)

    response = client.delete(
        f"/api/v1/budgets/{first_id}", headers=_headers(client, "gestor")
    )

    assert response.status_code == 204
    with app.state.session_factory() as db:
        assert db.get(BudgetEntry, first_id) is None


def test_read_only_cannot_delete_budget(budget_import_context):
    client, app, _ = budget_import_context
    first_id, _ = _seed_budgets(app)

    response = client.delete(
        f"/api/v1/budgets/{first_id}", headers=_headers(client, "consulta")
    )

    assert response.status_code == 403
    with app.state.session_factory() as db:
        assert db.get(BudgetEntry, first_id) is not None
        assert db.scalar(
            select(AuditLog).where(AuditLog.action == "BUDGET_DELETED")
        ) is None


def test_delete_unknown_budget_returns_404(budget_import_context):
    client, _, _ = budget_import_context

    response = client.delete(
        "/api/v1/budgets/999999", headers=_headers(client, "admin")
    )

    assert response.status_code == 404


def test_delete_budget_rolls_back_entry_and_audit_on_commit_failure(
    budget_import_context, monkeypatch,
):
    client, app, _ = budget_import_context
    first_id, second_id = _seed_budgets(app)
    headers = _headers(client, "admin")
    original_commit = Session.commit

    def fail_deletion_commit(session):
        if any(
            isinstance(item, AuditLog) and item.action == "BUDGET_DELETED"
            for item in session.new
        ):
            raise RuntimeError("falha controlada no commit")
        return original_commit(session)

    monkeypatch.setattr(Session, "commit", fail_deletion_commit)
    with pytest.raises(RuntimeError, match="falha controlada no commit"):
        client.delete(f"/api/v1/budgets/{first_id}", headers=headers)

    with app.state.session_factory() as db:
        assert db.get(BudgetEntry, first_id) is not None
        assert db.get(BudgetEntry, second_id) is not None
        assert db.scalar(
            select(AuditLog).where(AuditLog.action == "BUDGET_DELETED")
        ) is None


def test_validate_budget_returns_complete_preview_and_total(
    budget_import_context,
):
    client, _, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")

    response = _upload(
        client, "/api/v1/budgets/import/validate", path, _headers(client)
    )

    assert response.status_code == 200
    data = response.json()
    assert data["metadata"]["detected_type"] == "ORCAMENTO"
    assert data["metadata"]["year"] == 2026
    assert data["can_import"] is True
    assert data["totals"] == {"rows": 4, "value": 670.0}
    assert set(data["preview"][0]) == {
        "line", "year", "month", "entry_type", "category", "value",
        "source_label",
    }
    assert data["preview"][0]["source_label"] == "Repasse CDLs Estado Goiás"


def test_non_budget_workbook_is_rejected(budget_import_context):
    client, _, tmp_path = budget_import_context
    path = _target_workbook(tmp_path / "Metas 2026.xlsx")

    response = _upload(
        client, "/api/v1/budgets/import", path, _headers(client)
    )

    assert response.status_code == 422
    assert response.json()["detail"]["can_import"] is False


def test_internal_duplicate_blocks_entire_budget_import(budget_import_context):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx", duplicate=True)

    response = _upload(
        client, "/api/v1/budgets/import", path, _headers(client)
    )

    assert response.status_code == 422
    assert any(
        "registro repetido no arquivo" in error
        for error in response.json()["detail"]["errors"]
    )
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(BudgetEntry)) == 0


def test_persisted_duplicate_blocks_entire_budget_import(budget_import_context):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")
    with app.state.session_factory() as db:
        db.add(BudgetEntry(
            periodo_ano=2026, periodo_mes=1, tipo="RECEITA",
            categoria="RECEITA_DIRETA", descricao="Existente",
            valor_orcado=1,
        ))
        db.commit()

    response = _upload(
        client, "/api/v1/budgets/import", path, _headers(client)
    )

    assert response.status_code == 422
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(BudgetEntry)) == 1


@pytest.mark.parametrize("username", ["admin", "gestor"])
def test_authorized_roles_import_and_audit(budget_import_context, username):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / f"Orçamento {username} 2026.xlsx")

    response = _upload(
        client, "/api/v1/budgets/import", path, _headers(client, username)
    )

    assert response.status_code == 201
    assert response.json()["imported"] == 4
    with app.state.session_factory() as db:
        entries = list(db.scalars(select(BudgetEntry)))
        assert len(entries) == 4
        assert {
            (entry.tipo, entry.categoria)
            for entry in entries
        } == {
            ("RECEITA", "RECEITA_DIRETA"),
            ("DESPESA", "PESSOAL"),
        }
        assert all(
            entry.observacao == f"Importação operacional: {path.name}"
            for entry in entries
        )
        audit = db.scalar(select(AuditLog).where(
            AuditLog.action == "BUDGETS_IMPORTED"
        ))
        assert audit is not None
        assert audit.details == {
            "file_name": path.name,
            "year": 2026,
            "imported": 4,
            "inconsistencies": 0,
        }


def test_read_only_cannot_validate_or_import_budget(budget_import_context):
    client, _, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")
    headers = _headers(client, "consulta")

    assert _upload(
        client, "/api/v1/budgets/import/validate", path, headers
    ).status_code == 403
    assert _upload(
        client, "/api/v1/budgets/import", path, headers
    ).status_code == 403


def test_budget_import_rolls_back_batch_and_audit_on_failure(
    budget_import_context, monkeypatch,
):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")
    original = BudgetRepository.add_all

    def fail_batch(repository, entries):
        original(repository, entries)
        raise RuntimeError("falha controlada")

    monkeypatch.setattr(BudgetRepository, "add_all", fail_batch)
    with pytest.raises(RuntimeError, match="falha controlada"):
        _upload(client, "/api/v1/budgets/import", path, _headers(client))
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(BudgetEntry)) == 0
        assert db.scalar(select(AuditLog).where(
            AuditLog.action == "BUDGETS_IMPORTED"
        )) is None


def test_budget_validation_uses_one_existing_key_query(budget_import_context):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")
    statements = []

    def track(_connection, _cursor, statement, _parameters, _context, _many):
        if "budget_entries" in statement.casefold() and statement.lstrip().casefold().startswith("select"):
            statements.append(statement)

    event.listen(app.state.engine, "before_cursor_execute", track)
    try:
        response = _upload(
            client, "/api/v1/budgets/import/validate", path, _headers(client)
        )
    finally:
        event.remove(app.state.engine, "before_cursor_execute", track)

    assert response.status_code == 200
    assert len(statements) == 1


def test_budget_import_uses_add_all_without_individual_add(
    budget_import_context, monkeypatch,
):
    client, app, tmp_path = budget_import_context
    path = _budget_workbook(tmp_path / "Orçamento 2026.xlsx")
    calls = []
    original_add_all = BudgetRepository.add_all

    def reject_individual(*_args, **_kwargs):
        raise AssertionError("add individual não deve ser usado")

    def track_batch(repository, entries):
        materialized = list(entries)
        calls.append(len(materialized))
        return original_add_all(repository, materialized)

    monkeypatch.setattr(BudgetRepository, "add", reject_individual)
    monkeypatch.setattr(BudgetRepository, "add_all", track_batch)

    response = _upload(
        client, "/api/v1/budgets/import", path, _headers(client)
    )

    assert response.status_code == 201
    assert calls == [4]
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(BudgetEntry)) == 4
