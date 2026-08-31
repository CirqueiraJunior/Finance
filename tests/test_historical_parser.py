from decimal import Decimal
import inspect

from openpyxl import Workbook

import app.importers.historical_parser as parser_module
from app.importers.historical_importer import HistoricalWorkbookImporter
from app.importers.historical_parser import (
    BudgetImportData, HistoricalWorkbookParser, TargetImportData,
)
from app.models.entity import Entity
from app.repositories.entity_repository import EntityRepository
from app.services.entity_service import EntityService
from app.services.historical_import_service import HistoricalImportService


def _target_workbook(path) -> None:
    workbook = Workbook()
    target = workbook.active
    target.title = "Meta"
    target.append([])
    target.append([])
    target.append(["COD.", "Entidade", "JAN"])
    target.append([7500, "Consolidado", 999])
    target.append([7001, "Entidade inativa", 100])
    actual = workbook.create_sheet("Faturamento")
    actual.append([])
    actual.append([])
    actual.append(["COD.", "Entidade", "JAN"])
    actual.append([7500, "Consolidado", 999])
    actual.append([7001, "Entidade inativa", 80])
    workbook.save(path)


def _budget_workbook(path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planej. orçamentário"
    for _ in range(8):
        sheet.append([])
    for label, value in (
        ("Repasse CDLs Estado Goiás", 200),
        ("Despesas com Pessoal", 100),
        ("Despesas com Eventos", 50),
        ("Despesas com a Operação", 25),
        ("Software", 999),
        ("Viagem", 999),
        ("Impostos e Taxas", 999),
        ("Aplicação", 999),
        ("Resgate", 999),
        ("Saldo", 999),
    ):
        sheet.append([None, label, value, None])
    workbook.save(path)


def test_parser_is_pure_and_has_no_database_or_persistence_dependencies(tmp_path):
    path = tmp_path / "Meta x Realizado 2026.xlsx"
    _target_workbook(path)

    result = HistoricalWorkbookParser().parse(path)

    source = inspect.getsource(parser_module).casefold()
    assert "sqlalchemy" not in source
    assert "session" not in source
    assert "repository" not in source
    assert "commit(" not in source
    assert result.can_import
    assert isinstance(result.data[0], TargetImportData)


def test_target_parser_excludes_only_consolidated_code_and_preserves_values(tmp_path):
    path = tmp_path / "Meta x Realizado 2026.xlsx"
    _target_workbook(path)

    result = HistoricalWorkbookParser().parse(path)

    assert result.metadata.detected_type == "META_REALIZADO"
    assert result.metadata.year == 2026
    assert {(row.code, row.indicator, row.target, row.actual) for row in result.data} == {
        (7001, "CONSULTAS", Decimal("100.0000"), Decimal("80.0000"))
    }
    assert all(row.code != 7500 for row in result.data)
    assert result.warnings


def test_budget_parser_keeps_only_current_operational_categories(tmp_path):
    path = tmp_path / "Projeção Orçamentária 2026.xlsx"
    _budget_workbook(path)

    result = HistoricalWorkbookParser().parse(path)

    assert result.metadata.detected_type == "ORCAMENTO"
    assert all(isinstance(row, BudgetImportData) for row in result.data)
    assert {(row.entry_type, row.category) for row in result.data} == {
        ("RECEITA", "RECEITA_DIRETA"),
        ("DESPESA", "PESSOAL"),
        ("DESPESA", "EVENTOS"),
        ("DESPESA", "OPERACIONAL"),
    }
    assert not {"SOFTWARE", "VIAGEM", "IMPOSTOS_E_TAXAS", "SALDO_APLICADO"}.intersection(
        row.category for row in result.data
    )
    assert result.total == Decimal("375.0000")


def test_local_service_consumes_parser_adapter_and_accepts_inactive_entity(
        db_session, tmp_path):
    path = tmp_path / "Meta x Realizado 2026.xlsx"
    _target_workbook(path)
    db_session.add(Entity(
        codigo_entidade=7001, nome="Entidade inativa", ativa=False,
    ))
    db_session.commit()
    service = HistoricalImportService(
        db_session,
        HistoricalWorkbookImporter(HistoricalWorkbookParser()),
        EntityService(EntityRepository(db_session)),
        None,
        None,
        None,
    )

    preview = service.analyze(path)

    assert preview.detected_type == "META_REALIZADO"
    assert preview.errors == []
    assert preview.valid_rows == 1
    assert preview.rows[0]["entity_id"] is not None


def test_target_parser_separates_official_query_and_registration_sections(tmp_path):
    path = tmp_path / "Meta x Realizado - Oficial 2026.xlsm"
    workbook = Workbook()
    target = workbook.active
    target.title = "Meta"
    target.append(list(range(1, 15)))
    target.append(["META DE CONSULTAS"])
    target.append(["COD.", "Entidade", "JAN", "FEV"])
    target.append([7501, "Goiânia", 100, 110])
    target.append([])
    target.append(list(range(1, 15)))
    target.append(["META DE REGISTROS"])
    target.append(["COD.", "Entidade", "JAN", "FEV"])
    target.append([7501, "Goiânia", 20, 25])
    actual = workbook.create_sheet("Faturamento")
    actual.append(list(range(1, 15)))
    actual.append(["CONSULTAS REALIZADAS"])
    actual.append(["COD.", "Entidade", "JAN", "FEV"])
    actual.append([7501, "Goiânia", 80, 90])
    actual.append(list(range(1, 15)))
    actual.append(["REGISTROS REALIZADOS"])
    actual.append(["COD.", "Entidade", "JAN", "FEV"])
    actual.append([7501, "Goiânia", 15, 18])
    workbook.save(path)

    result = HistoricalWorkbookParser().parse(path)

    assert [
        (row.line, row.code, row.month, row.indicator, row.target, row.actual)
        for row in result.data
    ] == [
        (4, 7501, 1, "CONSULTAS", Decimal("100.0000"), Decimal("80.0000")),
        (4, 7501, 2, "CONSULTAS", Decimal("110.0000"), Decimal("90.0000")),
        (9, 7501, 1, "REGISTROS", Decimal("20.0000"), Decimal("15.0000")),
        (9, 7501, 2, "REGISTROS", Decimal("25.0000"), Decimal("18.0000")),
    ]
    assert all(row.code != 1 for row in result.data)
    assert not any("REGISTROS só serão importados" in item for item in result.warnings)
