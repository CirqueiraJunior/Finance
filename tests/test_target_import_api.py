from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import event, func, select

from app.models.entity import Entity
from app.models.target_entry import TargetEntry
from app.repositories.target_repository import TargetRepository
from finance_server.app_factory import create_app
from finance_server.config import ServerSettings
from finance_server.models import AuditLog, User, UserRole
from finance_server.security import hash_password


PASSWORD = "Strong!Pass123"


def _workbook(path: Path, *, two_months: bool = False) -> Path:
    workbook = Workbook()
    target = workbook.active
    target.title = "Meta"
    target.append([])
    target.append([])
    target.append(["COD.", "Entidade", "JAN", "FEV"])
    target.append([7500, "Consolidado", 999, 999 if two_months else None])
    target.append([7001, "Inativa", 100, 120 if two_months else None])
    actual = workbook.create_sheet("Faturamento")
    actual.append([])
    actual.append([])
    actual.append(["COD.", "Entidade", "JAN", "FEV"])
    actual.append([7500, "Consolidado", 999, 999 if two_months else None])
    actual.append([7001, "Inativa", 80, 90 if two_months else None])
    workbook.save(path)
    return path


def _workbook_with_both_indicators(path: Path) -> Path:
    workbook = Workbook()
    target = workbook.active
    target.title = "Meta"
    target.append(["META DE CONSULTAS"])
    target.append(["COD.", "Entidade", "JAN"])
    target.append([7001, "Inativa", 100])
    target.append(["META DE REGISTROS"])
    target.append(["COD.", "Entidade", "JAN"])
    target.append([7001, "Inativa", 20])
    actual = workbook.create_sheet("Faturamento")
    actual.append(["CONSULTAS REALIZADAS"])
    actual.append(["COD.", "Entidade", "JAN"])
    actual.append([7001, "Inativa", 80])
    actual.append(["REGISTROS REALIZADOS"])
    actual.append(["COD.", "Entidade", "JAN"])
    actual.append([7001, "Inativa", 15])
    workbook.save(path)
    return path


@pytest.fixture
def target_import_context(tmp_path):
    settings = ServerSettings(
        f"sqlite:///{(tmp_path / 'server.db').as_posix()}", "s" * 64
    )
    app = create_app(settings, create_schema=True)
    with app.state.session_factory() as db:
        db.add_all([
            User(nome="Admin", email="admin@test.local", username="admin",
                 password_hash=hash_password(PASSWORD),
                 perfil=UserRole.ADMINISTRATOR.value, ativo=True),
            User(nome="Gestor", email="gestor@test.local", username="gestor",
                 password_hash=hash_password(PASSWORD),
                 perfil=UserRole.MANAGER.value, ativo=True),
            User(nome="Consulta", email="consulta@test.local", username="consulta",
                 password_hash=hash_password(PASSWORD),
                 perfil=UserRole.READ_ONLY.value, ativo=True),
            Entity(codigo_entidade=7001, nome="Entidade inativa", ativa=False),
            Entity(codigo_entidade=7500, nome="Consolidado", ativa=True),
        ])
        db.commit()
    with TestClient(app) as client:
        yield client, app, tmp_path
    app.state.engine.dispose()


def _headers(client, username="admin"):
    pair = client.post("/api/v1/auth/login", json={
        "identifier": username, "password": PASSWORD,
    }).json()
    return {"Authorization": f"Bearer {pair['access_token']}"}


def _upload(client, endpoint, path, headers):
    with path.open("rb") as handle:
        return client.post(
            endpoint, headers=headers,
            files={"file": (path.name, handle)},
        )


def _seed_targets(app):
    with app.state.session_factory() as db:
        entity = db.scalar(select(Entity).where(Entity.codigo_entidade == 7001))
        first = TargetEntry(
            entity_id=entity.id, periodo_ano=2026, periodo_mes=8,
            indicador="CONSULTAS", valor_meta=100, valor_realizado=80,
            observacao="TESTE PRIMEIRO",
        )
        second = TargetEntry(
            entity_id=entity.id, periodo_ano=2026, periodo_mes=9,
            indicador="CONSULTAS", valor_meta=120, valor_realizado=90,
            observacao="TESTE SEGUNDO",
        )
        db.add_all([first, second])
        db.commit()
        return first.id, second.id


def test_administrator_deletes_only_requested_target_and_audits(
    target_import_context,
):
    client, app, _ = target_import_context
    first_id, second_id = _seed_targets(app)

    response = client.delete(
        f"/api/v1/targets/{first_id}", headers=_headers(client, "admin")
    )

    assert response.status_code == 204
    assert response.content == b""
    with app.state.session_factory() as db:
        assert db.get(TargetEntry, first_id) is None
        assert db.get(TargetEntry, second_id) is not None
        deletion = db.scalar(
            select(AuditLog).where(
                AuditLog.action == "TARGET_DELETED",
                AuditLog.entity_id == str(first_id),
            )
        )
        assert deletion is not None
        assert deletion.entity_type == "TargetEntry"
        assert deletion.details == {
            "entity_id": 1,
            "entity_code": 7001,
            "year": 2026,
            "month": 8,
            "indicator": "CONSULTAS",
            "target": "100.0000",
            "actual": "80.0000",
            "notes": "TESTE PRIMEIRO",
        }


def test_manager_can_delete_target(target_import_context):
    client, app, _ = target_import_context
    first_id, _ = _seed_targets(app)

    response = client.delete(
        f"/api/v1/targets/{first_id}", headers=_headers(client, "gestor")
    )

    assert response.status_code == 204
    with app.state.session_factory() as db:
        assert db.get(TargetEntry, first_id) is None


def test_read_only_cannot_delete_target(target_import_context):
    client, app, _ = target_import_context
    first_id, _ = _seed_targets(app)

    response = client.delete(
        f"/api/v1/targets/{first_id}", headers=_headers(client, "consulta")
    )

    assert response.status_code == 403
    with app.state.session_factory() as db:
        assert db.get(TargetEntry, first_id) is not None
        assert db.scalar(
            select(AuditLog).where(AuditLog.action == "TARGET_DELETED")
        ) is None


def test_delete_unknown_target_returns_404(target_import_context):
    client, _, _ = target_import_context

    response = client.delete(
        "/api/v1/targets/999999", headers=_headers(client, "admin")
    )

    assert response.status_code == 404


def test_validate_and_import_targets_accept_inactive_ignore_7500_and_audit(
        target_import_context):
    client, app, tmp_path = target_import_context
    path = _workbook(tmp_path / "Meta x Realizado 2026.xlsx")
    headers = _headers(client, "gestor")

    validation = _upload(
        client, "/api/v1/targets/import/validate", path, headers
    )
    assert validation.status_code == 200
    data = validation.json()
    assert data["metadata"]["detected_type"] == "META_REALIZADO"
    assert data["can_import"] is True
    assert data["totals"] == {"rows": 1, "target": 100.0, "actual": 80.0}
    assert data["preview"][0]["entity_code"] == 7001
    assert all(row["entity_code"] != 7500 for row in data["preview"])

    imported = _upload(client, "/api/v1/targets/import", path, headers)
    assert imported.status_code == 201
    assert imported.json()["imported"] == 1
    with app.state.session_factory() as db:
        entry = db.scalar(select(TargetEntry))
        assert entry is not None
        assert entry.entity.ativa is False
        audit = db.scalar(select(AuditLog).where(AuditLog.action == "TARGETS_IMPORTED"))
        assert audit is not None
        assert audit.details == {
            "file_name": path.name, "year": 2026,
            "imported": 1, "inconsistencies": 1,
        }


def test_import_supports_queries_and_registrations(target_import_context):
    client, app, tmp_path = target_import_context
    path = _workbook_with_both_indicators(
        tmp_path / "Meta x Realizado Oficial 2026.xlsx"
    )

    response = _upload(
        client, "/api/v1/targets/import", path, _headers(client, "gestor")
    )

    assert response.status_code == 201
    assert response.json()["imported"] == 2
    with app.state.session_factory() as db:
        entries = list(db.scalars(select(TargetEntry)))
        assert {
            (entry.indicador, entry.valor_meta, entry.valor_realizado)
            for entry in entries
        } == {
            ("CONSULTAS", 100, 80),
            ("REGISTROS", 20, 15),
        }


def test_target_import_revalidates_and_blocks_duplicates(target_import_context):
    client, _, tmp_path = target_import_context
    path = _workbook(tmp_path / "Meta x Realizado 2026.xlsx")
    headers = _headers(client)
    assert _upload(client, "/api/v1/targets/import", path, headers).status_code == 201

    validation = _upload(
        client, "/api/v1/targets/import/validate", path, headers
    ).json()
    assert validation["can_import"] is False
    assert any("já existe Meta" in error for error in validation["errors"])
    repeated = _upload(client, "/api/v1/targets/import", path, headers)
    assert repeated.status_code == 422


def test_read_only_cannot_validate_or_import_targets(target_import_context):
    client, _, tmp_path = target_import_context
    path = _workbook(tmp_path / "Meta x Realizado 2026.xlsx")
    headers = _headers(client, "consulta")

    assert _upload(
        client, "/api/v1/targets/import/validate", path, headers
    ).status_code == 403
    assert _upload(
        client, "/api/v1/targets/import", path, headers
    ).status_code == 403


def test_target_import_rolls_back_all_rows_and_audit_on_failure(
        target_import_context, monkeypatch):
    client, app, tmp_path = target_import_context
    path = _workbook(
        tmp_path / "Meta x Realizado 2026.xlsx", two_months=True
    )
    headers = _headers(client)
    original = TargetRepository.add_all

    def fail_batch(repository, targets):
        original(repository, targets)
        raise RuntimeError("falha controlada")

    monkeypatch.setattr(TargetRepository, "add_all", fail_batch)
    with pytest.raises(RuntimeError, match="falha controlada"):
        _upload(client, "/api/v1/targets/import", path, headers)
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(TargetEntry)) == 0
        assert db.scalar(select(AuditLog).where(
            AuditLog.action == "TARGETS_IMPORTED"
        )) is None


def test_validation_uses_single_existing_key_query(target_import_context):
    client, app, tmp_path = target_import_context
    path = _workbook(tmp_path / "Meta x Realizado 2026.xlsx", two_months=True)
    statements = []

    def track(_connection, _cursor, statement, _parameters, _context, _many):
        if "target_entries" in statement.casefold() and statement.lstrip().casefold().startswith("select"):
            statements.append(statement)

    event.listen(app.state.engine, "before_cursor_execute", track)
    try:
        response = _upload(
            client, "/api/v1/targets/import/validate", path, _headers(client)
        )
    finally:
        event.remove(app.state.engine, "before_cursor_execute", track)

    assert response.status_code == 200
    assert len(statements) == 1


def test_internal_duplicate_blocks_entire_import(target_import_context):
    client, app, tmp_path = target_import_context
    path = _workbook(tmp_path / "Meta x Realizado 2026.xlsx")
    workbook = load_workbook(path)
    workbook["Meta"].append([7001, "Inativa repetida", 100])
    workbook.save(path)

    response = _upload(
        client, "/api/v1/targets/import", path, _headers(client)
    )

    assert response.status_code == 422
    assert any(
        "registro repetido no arquivo" in message
        for message in response.json()["detail"]["errors"]
    )
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(TargetEntry)) == 0
        assert db.scalar(select(AuditLog).where(
            AuditLog.action == "TARGETS_IMPORTED"
        )) is None
