"""Parsing puro dos arquivos operacionais históricos de Metas e Orçamento."""

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join(
        "".join(char for char in text if not unicodedata.combining(char))
        .upper().strip().split()
    )


def decimal_value(value: object) -> Decimal:
    if isinstance(value, float):
        value = str(value)
    try:
        result = Decimal(value).quantize(Decimal("0.0001"))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ValueError("Valor numérico inválido.") from error
    if not result.is_finite() or result < 0:
        raise ValueError("Valor numérico inválido.")
    return result


@dataclass(frozen=True, slots=True)
class HistoricalParseMetadata:
    file_path: Path
    detected_type: str
    year: int | None
    sheets: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class TargetImportData:
    line: int
    code: int
    year: int
    month: int
    indicator: str
    target: Decimal
    actual: Decimal


@dataclass(frozen=True, slots=True)
class BudgetImportData:
    line: int
    year: int
    month: int
    entry_type: str
    category: str
    value: Decimal
    source_label: str


@dataclass(frozen=True, slots=True)
class HistoricalParseResult:
    metadata: HistoricalParseMetadata
    data: tuple[TargetImportData | BudgetImportData, ...] = ()
    warnings: tuple[str, ...] = ()
    errors: tuple[str, ...] = ()
    total: Decimal = Decimal("0.0000")

    @property
    def can_import(self) -> bool:
        return bool(self.data) and not self.errors


class HistoricalWorkbookParser:
    """Lê e valida workbooks sem conhecer banco, repositório ou persistência."""

    BUDGET_MAPPINGS = {
        "REPASSE CDLS ESTADO GOIAS": ("RECEITA", "RECEITA_DIRETA"),
        "DESPESAS COM PESSOAL": ("DESPESA", "PESSOAL"),
        "DESPESAS COM EVENTOS": ("DESPESA", "EVENTOS"),
        "DESPESAS COM A OPERACAO": ("DESPESA", "OPERACIONAL"),
    }

    def parse(self, file_path: str | Path) -> HistoricalParseResult:
        path = Path(file_path)
        if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.is_file():
            return self._unknown(path, (), "Arquivo Excel inválido ou inexistente.")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            names = {normalized(name): name for name in workbook.sheetnames}
            sheets = tuple(workbook.sheetnames)
            if "META" in names and "FATURAMENTO" in names:
                return self._targets(
                    path, sheets, workbook[names["META"]], workbook[names["FATURAMENTO"]],
                    has_association="ASSOCIACOES" in names,
                )
            if "PLANEJ. ORCAMENTARIO" in names:
                return self._budget(path, sheets, workbook[names["PLANEJ. ORCAMENTARIO"]])
            return self._unknown(path, sheets, "Estrutura operacional não reconhecida.")
        finally:
            workbook.close()

    def _targets(self, path: Path, sheets: tuple[str, ...], target_sheet,
                 actual_sheet, *, has_association: bool) -> HistoricalParseResult:
        year = self._year_from_name(path) or 2026
        actual_by_key: dict[tuple[str, int, int], object] = {}
        actual_rows, actual_indicators = self._indicator_rows(actual_sheet, {
            "CONSULTAS REALIZADAS": "CONSULTAS",
            "REGISTROS REALIZADOS": "REGISTROS",
        })
        for indicator, _line, code, values in actual_rows:
            for month in range(1, 13):
                actual_by_key[(indicator, code, month)] = (
                    values[month + 1] if len(values) > month + 1 else None
                )

        data: list[TargetImportData] = []
        errors: list[str] = []
        target_rows, target_indicators = self._indicator_rows(target_sheet, {
            "META DE CONSULTAS": "CONSULTAS",
            "META DE REGISTROS": "REGISTROS",
        })
        for indicator, line, code, values in target_rows:
            for month in range(1, 13):
                target = values[month + 1] if len(values) > month + 1 else None
                actual = actual_by_key.get((indicator, code, month))
                if target is None and actual is None:
                    continue
                try:
                    data.append(TargetImportData(
                        line, code, year, month, indicator,
                        decimal_value(target or 0), decimal_value(actual or 0),
                    ))
                except ValueError as error:
                    errors.append(f"Linha {line}, mês {month}: {error}")

        warnings = []
        if "REGISTROS" not in target_indicators or "REGISTROS" not in actual_indicators:
            warnings.append(
                "A estrutura analisada contém CONSULTAS. REGISTROS só serão importados quando houver aba oficial inequívoca."
            )
        if has_association:
            warnings.append(
                "A planilha também contém Associação; selecione esse tipo no preview para importá-la."
            )
        metadata = HistoricalParseMetadata(path, "META_REALIZADO", year, sheets)
        return HistoricalParseResult(metadata, tuple(data), tuple(warnings), tuple(errors))

    def _indicator_rows(self, sheet, labels: dict[str, str]):
        physical_rows = list(enumerate(
            sheet.iter_rows(min_row=1, values_only=True), start=1
        ))
        has_sections = any(
            normalized(values[0] if values else None) in labels
            for _, values in physical_rows
        )
        active_indicator = None if has_sections else "CONSULTAS"
        indicators: set[str] = set()
        rows = []
        for line, values in physical_rows:
            first = values[0] if values else None
            section = labels.get(normalized(first))
            if section is not None:
                active_indicator = section
                indicators.add(section)
                continue
            code = self._entity_code(first)
            entity_name = values[1] if len(values) > 1 else None
            if (active_indicator is None or code is None
                    or not isinstance(entity_name, str) or not entity_name.strip()):
                continue
            indicators.add(active_indicator)
            rows.append((active_indicator, line, code, values))
        return rows, indicators

    def _budget(self, path: Path, sheets: tuple[str, ...], sheet) -> HistoricalParseResult:
        year = self._year_from_name(path) or 2026
        data: list[BudgetImportData] = []
        warnings: list[str] = []
        errors: list[str] = []
        total = Decimal("0.0000")
        ignored = 0
        for line, values in enumerate(sheet.iter_rows(min_row=9, values_only=True), start=9):
            label = normalized(values[1] if len(values) > 1 else None)
            mapping = self.BUDGET_MAPPINGS.get(label)
            if mapping is None:
                if label and any(
                    values[index] for index in range(2, min(len(values), 26), 2)
                ):
                    ignored += 1
                continue
            for month in range(1, 13):
                index = 2 + (month - 1) * 2
                value = values[index] if len(values) > index else None
                if value is None:
                    continue
                try:
                    amount = decimal_value(value)
                except ValueError as error:
                    errors.append(f"Linha {line}, mês {month}: {error}")
                    continue
                data.append(BudgetImportData(
                    line, year, month, mapping[0], mapping[1], amount,
                    str(values[1] or "").strip(),
                ))
                total += amount
        if ignored:
            warnings.append(
                f"{ignored} linhas detalhadas sem correspondência inequívoca foram excluídas do preview."
            )
        metadata = HistoricalParseMetadata(path, "ORCAMENTO", year, sheets)
        return HistoricalParseResult(
            metadata, tuple(data), tuple(warnings), tuple(errors), total
        )

    @staticmethod
    def _entity_code(value: object) -> int | None:
        if not isinstance(value, (int, float)):
            return None
        code = int(value)
        return code if code > 0 and code != 7500 else None

    @staticmethod
    def _year_from_name(path: Path) -> int | None:
        match = re.search(r"\b(20\d{2})\b", path.name)
        return int(match.group(1)) if match else None

    @staticmethod
    def _unknown(path: Path, sheets: tuple[str, ...], error: str) -> HistoricalParseResult:
        return HistoricalParseResult(
            HistoricalParseMetadata(path, "DESCONHECIDO", None, sheets), errors=(error,)
        )
